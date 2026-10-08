"""PDF/DOCX parsers → ParsedDocument (spec §5).

Preserved per source: page numbers (PDF native; DOCX has no native
pagination so it is stored as a single logical page — see note below),
headings + levels, paragraphs, tables (serialized inline), original text,
and character offsets into full_text.

Table fidelity note: PDF table text is extracted inline in reading order
(true cell reconstruction needs pdfplumber/camelot — a later upgrade that
only touches this file). DOCX tables keep row/cell structure serialized
as pipe-delimited rows.
"""
from __future__ import annotations

import io
import re
import zipfile

from app.modules.document.errors import DocumentParseError
from app.modules.document.structures import Page, Paragraph, ParsedDocument, Section

PARSER_VERSION = "1.0.0"

_NUMBERED_RE = re.compile(r"^(?:section\s+)?(\d+(?:\.\d+)*)\s*([\.\)\:]?)\s*(.*)$", re.IGNORECASE)
_ARTICLE_HEADING = re.compile(r"^ARTICLE\s+([IVXLCM0-9]+)\s*[\-\:\.]?\s*(.{0,150})$", re.IGNORECASE)
_EXHIBIT_HEADING = re.compile(r"^(EXHIBIT|SCHEDULE|ANNEX)\s+([A-Z0-9]+)\s*[\-\:\.]?\s*(.{0,150})$", re.IGNORECASE)


def detect_heading(line: str) -> tuple[int, str | None, str] | None:
    """Return (level, section_number, heading_text) if the line is a heading.

    Bare leading numbers ("30 days notice…") are NOT headings: a numeric
    prefix only counts with a dot ("8.2"), a separator ("1.", "2)", "3:"),
    or the word "Section". This avoids list items / quantities splitting clauses.
    """
    s = line.strip()
    if not s or len(s) > 200:
        return None
    m = _NUMBERED_RE.match(s)
    if m:
        num, sep, rest = m.group(1), m.group(2), m.group(3)
        if sep or "." in num or s.lower().startswith("section"):
            if rest.strip() or sep or "." in num or s.lower().startswith("section"):
                return (num.count(".") + 1, num, s)
            return None
    m = _ARTICLE_HEADING.match(s)
    if m:
        return (1, f"Article {m.group(1).upper()}", s)
    m = _EXHIBIT_HEADING.match(s)
    if m:
        return (1, f"{m.group(1).capitalize()} {m.group(2).upper()}", s)
    # Unnumbered ALL-CAPS headings, e.g. "LIMITATION OF LIABILITY".
    letters = [c for c in s if c.isalpha()]
    if len(s) <= 120 and len(letters) >= 3 and sum(c.isupper() for c in letters) / len(letters) > 0.85:
        if not s.endswith((".", ";", ",")):
            return (1, None, s)
    return None


def _assemble(pages_text: list[str]) -> tuple[str, list[tuple[int, int]]]:
    """Join page texts; return (full_text, per-page (char_start, char_end))."""
    parts: list[str] = []
    spans: list[tuple[int, int]] = []
    pos = 0
    for i, t in enumerate(pages_text):
        if i > 0:
            parts.append("\n\n")
            pos += 2
        spans.append((pos, pos + len(t)))
        parts.append(t)
        pos += len(t)
    return "".join(parts), spans


def _split_paragraphs(page_text: str, page_number: int, base: int) -> list[Paragraph]:
    """Blank-line separated paragraphs with offsets relative to full_text."""
    paras: list[Paragraph] = []
    # Split into blocks, then locate each block to get exact offsets.
    blocks = [b for b in re.split(r"\n\s*\n", page_text) if b.strip()]
    cursor = 0
    for block in blocks:
        idx = page_text.find(block, cursor)
        if idx == -1:  # defensive: fall back to cursor
            idx = cursor
        cursor = idx + len(block)
        text = block.strip()
        start = base + idx + (len(block) - len(block.lstrip()))
        end = base + idx + len(block.rstrip())
        paras.extend(_block_to_paras(text, page_number, start, end))
    return paras


def _block_to_paras(text: str, page_number: int, start: int, end: int) -> list[Paragraph]:
    """Split a block at every heading line (headings can sit mid-page when the
    PDF has no blank lines); other lines stay grouped as one body paragraph."""
    lines = text.split("\n")
    out: list[Paragraph] = []
    body: list[str] = []
    body_start: int | None = None
    pos = start

    def flush_body(upto: int) -> None:
        nonlocal body, body_start
        if body and body_start is not None:
            joined = "\n".join(body).strip()
            if joined:
                out.append(Paragraph(joined, page_number, body_start, upto))
        body, body_start = [], None

    for ln in lines:
        stripped = ln.strip()
        hit = detect_heading(stripped) if stripped else None
        lead = len(ln) - len(ln.lstrip())
        if hit:
            flush_body(pos)
            out.append(Paragraph(stripped, page_number, pos + lead, pos + len(ln.rstrip()),
                                 is_heading=True, heading_level=hit[0], section_number=hit[1]))
        else:
            if body_start is None:
                body_start = pos + (lead if stripped else 0)
            body.append(ln)
        pos += len(ln) + 1  # +1 for the "\n" join
    flush_body(end)
    return out or [Paragraph(text, page_number, start, end)]


def _sections_from_paras(paras: list[Paragraph], pages: list[Page]) -> list[Section]:
    sections: list[Section] = []
    heads = [p for p in paras if p.is_heading]
    for i, h in enumerate(heads):
        cend = heads[i + 1].char_start if i + 1 < len(heads) else (paras[-1].char_end if paras else h.char_end)
        sections.append(Section(
            heading=h.text[:200], level=h.heading_level, order_index=i,
            section_number=h.section_number, char_start=h.char_start, char_end=cend,
        ))
    for s in sections:
        s.page_start = _page_for_offset(pages, s.char_start)
        s.page_end = _page_for_offset(pages, max(s.char_end - 1, s.char_start))
    return sections


def _page_for_offset(pages: list[Page], offset: int) -> int | None:
    for p in pages:
        if p.char_start <= offset < p.char_end or (offset == p.char_end and p.char_end > p.char_start):
            return p.page_number
    return pages[-1].page_number if pages else None


# ---------------------------------------------------------------- PDF ----


def parse_pdf(data: bytes) -> ParsedDocument:
    try:
        from pypdf import PdfReader
    except ImportError as e:  # pragma: no cover
        raise DocumentParseError("PDF parser dependency missing (pypdf).") from e
    try:
        reader = PdfReader(io.BytesIO(data))
    except Exception as e:
        raise DocumentParseError(f"Could not read PDF (file may be corrupt): {e}") from e
    if reader.is_encrypted:
        try:
            if reader.decrypt("") == 0:
                raise DocumentParseError("PDF is password-protected; unlock it before uploading.")
        except DocumentParseError:
            raise
        except Exception as e:
            raise DocumentParseError(f"PDF is encrypted and could not be opened: {e}") from e
    if not reader.pages:
        raise DocumentParseError("PDF has no pages.")
    try:
        texts = [pg.extract_text() or "" for pg in reader.pages]
    except Exception as e:
        raise DocumentParseError(f"PDF text extraction failed (file may be corrupt): {e}") from e
    if not any(t.strip() for t in texts):
        raise DocumentParseError(
            "No extractable text found. This looks like a scanned/image-only PDF; "
            "OCR support is planned for a later phase."
        )
    full, spans = _assemble(texts)
    pages = [Page(i + 1, t, cs, ce) for i, (t, (cs, ce)) in enumerate(zip(texts, spans))]
    paras: list[Paragraph] = []
    for pg, (cs, _ce) in zip(pages, spans):
        for p in _split_paragraphs(pg.text, pg.page_number, cs):
            paras.append(p)
            pg.paragraph_spans.append((p.char_start, p.char_end))
    return ParsedDocument(full, pages, _sections_from_paras(paras, pages), paras,
                          "pypdf", PARSER_VERSION, len(pages))


# --------------------------------------------------------------- DOCX ----


def parse_docx(data: bytes) -> ParsedDocument:
    try:
        import docx
    except ImportError as e:  # pragma: no cover
        raise DocumentParseError("DOCX parser dependency missing (python-docx).") from e
    try:
        doc = docx.Document(io.BytesIO(data))
    except (zipfile.BadZipFile, ValueError) as e:
        raise DocumentParseError(f"Could not read DOCX (file may be corrupt): {e}") from e
    except Exception as e:
        raise DocumentParseError(f"Could not read DOCX: {e}") from e

    # NOTE: DOCX has no native pagination, so the whole document is stored as
    # one logical page (page 1). Order, headings, tables and offsets are real.
    raw: list[tuple[str, str | None]] = []  # (text, style_name); tables as (text, None)
    body = doc.element.body
    by_id = {id(p._p): p for p in doc.paragraphs}
    for child in body:
        tag = child.tag.split("}")[-1]
        if tag == "p":
            p = by_id.get(id(child))
            if p is None or not p.text.strip():
                continue
            raw.append((p.text.strip(), p.style.name if p.style else None))
        elif tag == "tbl":
            rows = []
            for tr in child.findall(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tr"):
                cells = []
                for tc in tr.findall("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tc"):
                    cell_text = " ".join(
                        t.text or "" for t in
                        tc.findall(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t")
                    ).strip()
                    cells.append(cell_text)
                if any(cells):
                    rows.append("| " + " | ".join(cells) + " |")
            if rows:
                raw.append(("\n".join(rows), None))

    if not raw:
        raise DocumentParseError("No extractable text found in this DOCX file.")
    full, _ = _assemble(["\n\n".join(t for t, _ in raw)])
    paras: list[Paragraph] = []
    cursor = 0
    for text, style in raw:
        idx = full.find(text, cursor)
        start = idx if idx != -1 else cursor
        end = start + len(text)
        cursor = end
        is_head, level, num = False, 1, None
        if style and style.lower().startswith("heading"):
            try:
                level = int(style.split()[-1])
            except ValueError:
                level = 1
            is_head, num = True, None
            hit = detect_heading(text.split("\n")[0])
            if hit:
                num = hit[1]
        else:
            hit = detect_heading(text.split("\n")[0])
            if hit and len(text) <= 400:
                is_head, (level, num, _) = True, hit
        paras.append(Paragraph(text, 1, start, end, is_heading=is_head,
                               heading_level=level, section_number=num))
    pages = [Page(1, full, 0, len(full), [(p.char_start, p.char_end) for p in paras])]
    return ParsedDocument(full, pages, _sections_from_paras(paras, pages), paras,
                          "python-docx", PARSER_VERSION, 1)


def parse_document(filename: str, data: bytes) -> ParsedDocument:
    """Dispatch by extension. Raises DocumentParseError for anything else."""
    lower = filename.lower()
    if lower.endswith(".pdf"):
        return parse_pdf(data)
    if lower.endswith(".docx"):
        return parse_docx(data)
    raise DocumentParseError(f"Unsupported file type for parsing: {filename!r} (expected .pdf/.docx).")
