"""In-test document builders (reportlab / python-docx). Real files, real parsers."""
import io


def make_pdf(pages: list[list[str]]) -> bytes:
    """pages → list of pages; each page is a list of text lines."""
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=letter)
    for lines in pages:
        y = 750
        for line in lines:
            if line:
                c.drawString(72, y, line)
            y -= 14
            if y < 72:
                c.showPage()
                y = 750
        c.showPage()
    c.save()
    return buf.getvalue()


def make_blank_pdf(n_pages: int = 2) -> bytes:
    """Image-only-scan stand-in: pages with no extractable text."""
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=letter)
    for _ in range(n_pages):
        c.showPage()
    c.save()
    return buf.getvalue()


def make_docx(blocks: list[tuple[str, object]]) -> bytes:
    """blocks: ("h1"|"h2"|"p", text) or ("table", [[cell, ...], ...])."""
    import docx

    doc = docx.Document()
    for kind, payload in blocks:
        if kind == "h1":
            doc.add_heading(str(payload), level=1)
        elif kind == "h2":
            doc.add_heading(str(payload), level=2)
        elif kind == "p":
            doc.add_paragraph(str(payload))
        elif kind == "table":
            rows = payload
            assert isinstance(rows, list)
            t = doc.add_table(rows=len(rows), cols=len(rows[0]))
            for i, row in enumerate(rows):
                for j, val in enumerate(row):
                    t.cell(i, j).text = val
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


CONTRACT_PDF_PAGES = [
    ["1. Payment Terms",
     "The Client shall pay the Vendor within 30 days of invoice.",
     "30 days is the standard payment term.",  # bare number: must NOT split
     "Late payments accrue interest at 1% per month."],
    ["2. Confidentiality",
     "Each party shall keep confidential information strictly confidential.",
     "This obligation survives termination by two years."],
    ["8.2 Limitation of Liability",
     "Neither party shall be liable for indirect damages.",
     "Total liability is capped at fees paid in 12 months."],
]

CONTRACT_DOCX_BLOCKS: list[tuple[str, object]] = [
    ("h1", "1. Payment Terms"),
    ("p", "The Client shall pay the Vendor within 45 days of invoice."),
    ("h1", "2. Termination"),
    ("p", "Either party may terminate with 30 days written notice."),
    ("p", "8.2.1 Subcontractor liability remains with the Vendor at all times."),
    ("table", [["Item", "Amount"], ["Fees", "$10,000"]]),
]
